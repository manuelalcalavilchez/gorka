"""Generación de Excel, PDF y ZIP de un estudio."""
from __future__ import annotations

import json
import zipfile
from datetime import datetime
from pathlib import Path

import geopandas as gpd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (CondPageBreak, HRFlowable, Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)
from shapely.geometry import shape

from . import render, settings, terreno


def _registrar_fuentes() -> None:
    """DejaVu Sans (incluida con matplotlib) para tener ², ³, € y acentos en el PDF."""
    from matplotlib import font_manager
    from matplotlib.font_manager import FontProperties
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVu" in pdfmetrics.getRegisteredFontNames():
        return
    def f(**kw):
        return font_manager.findfont(FontProperties(family="DejaVu Sans", **kw), fallback_to_default=True)
    pdfmetrics.registerFont(TTFont("DejaVu", f()))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", f(weight="bold")))
    pdfmetrics.registerFont(TTFont("DejaVu-Oblique", f(style="oblique")))
    pdfmetrics.registerFont(TTFont("DejaVu-BoldOblique", f(weight="bold", style="oblique")))
    addMapping("DejaVu", 0, 0, "DejaVu")
    addMapping("DejaVu", 1, 0, "DejaVu-Bold")
    addMapping("DejaVu", 0, 1, "DejaVu-Oblique")
    addMapping("DejaVu", 1, 1, "DejaVu-BoldOblique")
from .estudio import carpeta, resultado

AZUL = colors.HexColor("#1F2A44")
ACENTO = colors.HexColor("#0F7B6C")
SEM_COLOR = {"rojo": colors.HexColor("#C62828"), "ambar": colors.HexColor("#EF8F00"),
             "verde": colors.HexColor("#2E7D32"), "info": colors.HexColor("#1E88E5"), "gris": colors.HexColor("#9E9E9E")}
SEM_TEXTO = {"rojo": "Excluyente", "ambar": "Condicionante", "verde": "Sin afección", "info": "Informativo",
             "gris": "Sin datos"}
ESTADO_TEXTO = {"afecta": "Afecta", "sin_afeccion": "Sin afección", "no_instalada": "Capa no instalada",
                "sin_cobertura": "Sin cobertura en la comunidad", "error": "Error"}


def _n(v, dec=0):
    if v is None:
        return "—"
    s = f"{v:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# ================================================================== EXCEL
def excel(r: dict, ruta: Path) -> None:
    wb = Workbook()
    cab = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F2A44")

    def hoja(titulo, filas, encabezado):
        ws = wb.create_sheet(titulo[:31])
        ws.append(encabezado)
        for c in ws[1]:
            c.font, c.fill = cab, relleno
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for f in filas:
            ws.append(list(f))
        for i, _ in enumerate(encabezado, 1):
            ancho = max(len(str(ws.cell(row=j, column=i).value or "")) for j in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 70)
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    g, a = r["global"], r["ambito"]
    filas = [("Estudio", r["nombre"]), ("Modo", r["modo_nombre"]), ("Fecha", r["fecha"]),
             ("Semáforo global", SEM_TEXTO.get(g["semaforo"], g["semaforo"])), ("Valoración", g["texto"]),
             ("Comunidad(es)", ", ".join(c["nombre"] for c in r["territorio"].get("ccaa", []))),
             ("Municipio(s)", ", ".join(m["nombre"] for m in r["territorio"].get("municipios", []))),
             ("Fichero del ámbito", a["fichero"]), ("Sistema de referencia", a["crs"]),
             ("Superficie del ámbito (m²)", a["area_m2"]), ("Superficie de obra (m²)", a["area_obra_m2"]),
             ("Longitud (m)", a["longitud_m"] if a["tipo"] == "linea" else None),
             ("MDT", (r["terreno"]["fuente"] or {}).get("nombre", "Sin MDT")),
             ("MDT fiable para volúmenes", "Sí" if r["terreno"]["fiable"] else "No")]
    filas += [("Aviso", av) for av in r.get("avisos", [])]
    filas.append(("Emitido por", f"{settings.MARCA} · {settings.MARCA_WEB}"))
    hoja("Resumen", filas, ["Concepto", "Valor"])

    st = r["terreno"].get("estadisticas") or {}
    if st and not st.get("sin_datos"):
        t = [(k.replace("_", " ").capitalize(), v) for k, v in st.items() if not isinstance(v, list)]
        t += [(f"Pendiente {c['clase']} (%)", c["pct"]) for c in st.get("clases_pendiente", [])]
        hoja("Terreno", t, ["Indicador", "Valor"])

    if r.get("volumenes"):
        v = [(k.replace("_", " ").capitalize(), val) for k, val in r["volumenes"].items()]
        v += [(f"Parámetro: {k.replace('_', ' ')}", val) for k, val in r["parametros"].items()]
        hoja("Movimiento de tierras", v, ["Concepto", "Valor"])
    if r.get("vial"):
        hoja("Tramos vial", [(t["pk_ini"], t["pk_fin"], t["desmonte_m3"], t["terraplen_m3"],
                              t["pendiente_media_pct"], t["altura_max_m"]) for t in r["vial"]["tramos"]],
             ["PK inicio", "PK fin", "Desmonte (m³)", "Terraplén (m³)", "Pendiente media (%)", "Altura máx. (m)"])

    filas = []
    for af in r["afecciones"]:
        if af.get("zonas"):
            for z in af["zonas"]:
                filas.append((af["tema"], af["nombre"], z["etiqueta"], SEM_TEXTO.get(z["semaforo"]),
                              "Sí" if z["afecta"] else "No", z["area_m2"], z["pct"],
                              ", ".join(z.get("elementos", [])), af.get("distancia_min_m")))
        else:
            filas.append((af["tema"], af["nombre"], ESTADO_TEXTO.get(af["estado"], af["estado"]),
                          SEM_TEXTO.get(af["semaforo"]), "No", 0, 0, "", af.get("distancia_min_m")))
    hoja("Afecciones", filas, ["Tema", "Capa", "Zona", "Clasificación", "Afecta", "Superficie (m²)", "% obra",
                               "Elementos", "Distancia mínima (m)"])
    hoja("Tramitación", [(t["nivel"], t["tramite"], t["motivo"]) for t in r["tramitacion"]],
         ["Prioridad", "Trámite", "Motivo"])
    hoja("Fuentes", [(af["nombre"], af["fuente"], af["fecha"], ESTADO_TEXTO.get(af["estado"], af["estado"]),
                      "Sí" if af.get("cobertura_parcial") else "No", af.get("notas", ""))
                     for af in r["afecciones"]]
         + ([(f"MDT: {r['terreno']['fuente']['nombre']}", r["terreno"]["fuente"]["fuente"], "", "Usado", "",
              ", ".join(r["terreno"]["fuente"]["ficheros"]))] if r["terreno"]["fuente"] else []),
         ["Capa", "Fuente", "Fecha/versión", "Estado", "Cobertura parcial", "Notas"])
    for ws in wb.worksheets:
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(ruta)


# ================================================================== PLANOS
def _planos(r: dict, d: Path, dest: Path) -> dict:
    out = {}
    amb = gpd.read_file(d / "ambito.geojson").to_crs(r["ambito"]["crs"]).geometry.iloc[0]
    crs = r["ambito"]["crs"]
    geom_ref = amb
    if (d / "visor/ocupacion.geojson").exists():
        geom_ref = gpd.read_file(d / "visor/ocupacion.geojson").to_crs(crs).geometry.iloc[0].union(amb)
    zonas = []
    for af in r["afecciones"]:
        f = d / "capas" / f"{af['id']}.geojson"
        if af.get("estado") != "afecta" or not f.exists():
            continue
        gj = json.loads(f.read_text(encoding="utf-8"))
        feats = [x for x in gj["features"] if x["properties"].get("tipo") == "zona"]
        if feats:
            gs = gpd.GeoSeries([shape(x["geometry"]) for x in feats], crs="EPSG:4326").to_crs(crs)
            zonas += [(g, x["properties"]["semaforo"], x["properties"]["etiqueta"]) for g, x in zip(gs, feats)]
    render.plano_afecciones(geom_ref, zonas, crs, dest / "plano_afecciones.jpg")
    out["afecciones"] = dest / "plano_afecciones.jpg"
    if (d / "terreno.npz").exists():
        t = terreno.Terreno.cargar(d / "terreno.npz")
        p = terreno.pendiente_pct(t.z, t.res)
        p_vis = np.where(t.mascara, p, np.nan) if r["ambito"]["tipo"] == "poligono" else p
        render.plano_raster(geom_ref, p_vis, t.transform, crs, dest / "plano_pendientes.jpg", "pendientes")
        out["pendientes"] = dest / "plano_pendientes.jpg"
        st = r["terreno"].get("estadisticas") or {}
        if st.get("clases_pendiente"):
            render.grafico_pendientes(st["clases_pendiente"], dest / "grafico_pendientes.png")
            out["grafico_pendientes"] = dest / "grafico_pendientes.png"
        if (d / "obra.npz").exists():
            dif = np.load(d / "obra.npz")["dif"]
            render.plano_raster(geom_ref, dif, t.transform, crs, dest / "plano_corte_relleno.jpg", "corte")
            out["corte"] = dest / "plano_corte_relleno.jpg"
    if r.get("vial"):
        render.grafico_perfil(r["vial"]["perfil"], dest / "perfil_longitudinal.png")
        out["perfil"] = dest / "perfil_longitudinal.png"
    return out


# ================================================================== PDF
class _Lienzo(rl_canvas.Canvas):
    def __init__(self, *a, titulo="", **k):
        super().__init__(*a, **k)
        self._paginas = []
        self._titulo = titulo

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for est in self._paginas:
            self.__dict__.update(est)
            self._pie(total)
            super().showPage()
        super().save()

    def _pie(self, total):
        w, _ = A4
        self.saveState()
        self.setStrokeColor(ACENTO)
        self.setLineWidth(0.8)
        self.line(18 * mm, 14 * mm, w - 18 * mm, 14 * mm)
        self.setFont("DejaVu", 7)
        self.setFillColor(colors.HexColor("#555555"))
        self.drawString(18 * mm, 10 * mm, f"{settings.MARCA} · {settings.MARCA_WEB} · {self._titulo}")
        self.drawRightString(w - 18 * mm, 10 * mm, f"Página {self._pageNumber} de {total}")
        self.restoreState()


def _img(ruta: Path, ancho_max=170 * mm, alto_max=150 * mm):
    with PILImage.open(ruta) as im:
        w, h = im.size
    esc = min(ancho_max / w, alto_max / h)
    return Image(str(ruta), width=w * esc, height=h * esc)


def _tabla(datos, anchos, cab=True, estilo_extra=None):
    t = Table(datos, colWidths=anchos, repeatRows=1 if cab else 0)
    est = [("FONT", (0, 0), (-1, -1), "DejaVu", 8), ("VALIGN", (0, 0), (-1, -1), "TOP"),
           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D5DD")),
           ("ROWBACKGROUNDS", (0, 1 if cab else 0), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    if cab:
        est += [("BACKGROUND", (0, 0), (-1, 0), AZUL), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONT", (0, 0), (-1, 0), "DejaVu-Bold", 8)]
    t.setStyle(TableStyle(est + (estilo_extra or [])))
    return t


def pdf(r: dict, ruta: Path, planos: dict) -> None:
    _registrar_fuentes()
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="DejaVu-Bold", fontSize=13, textColor=AZUL,
                        spaceBefore=6, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="DejaVu-Bold", fontSize=10.5, textColor=ACENTO,
                        spaceBefore=8, spaceAfter=4)
    cuerpo = ParagraphStyle("c", parent=ss["Normal"], fontName="DejaVu", fontSize=8.8, leading=12)
    peq = ParagraphStyle("p", parent=cuerpo, fontSize=7.5, leading=10, textColor=colors.HexColor("#444"))
    celda = ParagraphStyle("cel", parent=cuerpo, fontSize=7.8, leading=9.6)

    doc = SimpleDocTemplate(str(ruta), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=20 * mm, title=f"Prefactibilidad — {r['nombre']}",
                            author=f"{settings.MARCA} ({settings.MARCA_WEB})")
    s = []
    g, a, st = r["global"], r["ambito"], r["terreno"].get("estadisticas") or {}

    # ---------- portada / cabecera
    s.append(Paragraph(f"<font color='#0F7B6C'><b>{settings.MARCA.upper()}</b></font>"
                       f"<font size=8 color='#555'>  ·  {settings.MARCA_WEB}</font>", cuerpo))
    s.append(Spacer(1, 4))
    s.append(Paragraph("ESTUDIO DE PREFACTIBILIDAD DE MOVIMIENTO DE TIERRAS", h1))
    s.append(Paragraph(f"<b>{r['nombre']}</b> — {r['modo_nombre']}", cuerpo))
    s.append(Paragraph(f"Fecha: {datetime.fromisoformat(r['fecha']).strftime('%d/%m/%Y %H:%M')} · "
                       f"Referencia: {r['id'][:8].upper()}-v{r['version']}", peq))
    s.append(HRFlowable(width="100%", thickness=1.2, color=ACENTO, spaceBefore=6, spaceAfter=8))

    # ---------- 1. resumen
    s.append(Paragraph("1. Resumen ejecutivo", h2))
    sem = g["semaforo"]
    fila_sem = [[Paragraph(f"<font color='white'><b>{SEM_TEXTO[sem].upper()}</b></font>", celda),
                 Paragraph(g["texto"], celda)]]
    s.append(_tabla(fila_sem, [35 * mm, 139 * mm], cab=False,
                    estilo_extra=[("BACKGROUND", (0, 0), (0, 0), SEM_COLOR[sem]), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    s.append(Spacer(1, 5))
    vol = r.get("volumenes") or {}
    datos = [["Indicador", "Valor"],
             ["Comunidad / municipio", Paragraph(", ".join(c["nombre"] for c in r["territorio"].get("ccaa", [])) + " · "
                                                 + ", ".join(m["nombre"] for m in r["territorio"].get("municipios", [])[:4]), celda)],
             ["Superficie del ámbito", f"{_n(a['area_m2'])} m²" if a["tipo"] == "poligono" else f"Eje de {_n(a['longitud_m'])} m"],
             ["Superficie de obra (con taludes)", f"{_n(a['area_obra_m2'])} m²"]]
    if vol:
        datos += [["Desmonte / terraplén", f"{_n(vol.get('volumen_desmonte_m3'))} / {_n(vol.get('volumen_terraplen_m3'))} m³"],
                  ["Balance (desmonte − terraplén)", f"{_n(vol.get('balance_m3'))} m³"]]
    if st and not st.get("sin_datos"):
        datos += [["Cotas (mín. – máx.)", f"{_n(st['cota_min'], 1)} – {_n(st['cota_max'], 1)} m"],
                  ["Pendiente media / P90", f"{_n(st['pendiente_media'], 1)} % / {_n(st['pendiente_p90'], 1)} %"]]
    if g["motivos"]:
        datos.append(["Condicionantes principales", Paragraph(", ".join(g["motivos"]), celda)])
    s.append(_tabla(datos, [60 * mm, 114 * mm]))
    if r.get("avisos") or g.get("capas_sin_datos"):
        s.append(Spacer(1, 4))
        for av in r.get("avisos", []):
            s.append(Paragraph(f"<b>Aviso:</b> {av}", peq))
        if g.get("capas_sin_datos"):
            s.append(Paragraph("Capas sin datos para este ámbito: " + ", ".join(g["capas_sin_datos"]) + ".", peq))
        if g.get("capas_cobertura_parcial"):
            s.append(Paragraph("Capas con cobertura parcial (la ausencia de afección no es concluyente): "
                               + ", ".join(g["capas_cobertura_parcial"]) + ".", peq))

    # ---------- 2. ámbito
    s.append(CondPageBreak(40 * mm))
    s.append(Paragraph("2. Ámbito y localización", h2))
    terr = r["territorio"]
    datos = [["Dato", "Valor"], ["Fichero", a["fichero"]], ["Sistema de referencia", a["crs"]],
             ["Centroide (lon, lat)", f"{a['centroide'][0]:.5f}, {a['centroide'][1]:.5f}"],
             ["Municipios", Paragraph(", ".join(f"{m['nombre']} ({_n(m['pct'], 1)} %)" for m in terr.get("municipios", [])) or "—", celda)]]
    s.append(_tabla(datos, [60 * mm, 114 * mm]))
    if "afecciones" in planos:
        s.append(Spacer(1, 6))
        s.append(KeepTogether([_img(planos["afecciones"])]))

    # ---------- 3. terreno
    s.append(PageBreak())
    s.append(Paragraph("3. Terreno", h2))
    f = r["terreno"]["fuente"]
    if f:
        s.append(Paragraph(f"Modelo del terreno: <b>{f['nombre']}</b> (paso de trabajo {_n(f['res_trabajo_m'], 1)} m). "
                           f"Fuente: {f['fuente']}.", cuerpo))
        if not r["terreno"]["fiable"]:
            s.append(Paragraph("<font color='#C62828'><b>Resolución insuficiente: pendientes y volúmenes solo son "
                               "orientativos.</b></font>", cuerpo))
    else:
        s.append(Paragraph("No hay MDT disponible para el ámbito.", cuerpo))
    if st and not st.get("sin_datos"):
        datos = [["Indicador", "Valor"], ["Cota mínima / media / máxima",
                 f"{_n(st['cota_min'], 1)} / {_n(st['cota_media'], 1)} / {_n(st['cota_max'], 1)} m"],
                 ["Desnivel", f"{_n(st['desnivel'], 1)} m"],
                 ["Pendiente media / P90 / máxima",
                  f"{_n(st['pendiente_media'], 1)} / {_n(st['pendiente_p90'], 1)} / {_n(st['pendiente_max'], 1)} %"]]
        datos += [[f"Superficie con pendiente {c['clase']}", f"{_n(c['pct'], 1)} % ({_n(c['area_m2'])} m²)"]
                  for c in st["clases_pendiente"]]
        s.append(Spacer(1, 4))
        s.append(_tabla(datos, [70 * mm, 104 * mm]))
    if "grafico_pendientes" in planos:
        s.append(_img(planos["grafico_pendientes"], 150 * mm, 60 * mm))
    if "pendientes" in planos:
        s.append(KeepTogether([_img(planos["pendientes"], 170 * mm, 130 * mm)]))

    # ---------- 4. movimiento de tierras
    if vol:
        s.append(PageBreak())
        s.append(Paragraph("4. Movimiento de tierras", h2))
        etiquetas = {
            "cota_plataforma": "Cota de plataforma (m)", "cota_compensada": "Cota calculada por compensación",
            "cota_final": "Cota final de relleno (m)", "capacidad_m3": "Capacidad de relleno (m³)",
            "cota_fondo": "Cota de fondo (m)", "profundidad_max_m": "Profundidad máxima (m)",
            "volumen_extraible_m3": "Volumen extraíble en banco (m³)", "longitud_m": "Longitud del eje (m)",
            "volumen_desmonte_m3": "Volumen de desmonte (m³)", "volumen_terraplen_m3": "Volumen de terraplén (m³)",
            "balance_m3": "Balance desmonte − terraplén (m³)", "excedente_esponjado_m3": "Excedente esponjado (m³)",
            "deficit_m3": "Déficit de tierras (m³)", "viajes_camion": "Viajes de camión (excedente)",
            "viajes_camion_aporte": "Viajes de camión (aporte)",
            "altura_max_desmonte_m": "Altura máx. de desmonte (m)", "altura_max_terraplen_m": "Altura máx. de terraplén (m)",
            "superficie_ocupada_m2": "Superficie ocupada (m²)", "superficie_taludes_m2": "Superficie de taludes (m²)",
            "pendiente_max_rasante_pct": "Pendiente máx. de rasante (%)",
            "pendiente_max_terreno_pct": "Pendiente máx. del terreno en el eje (%)",
            "ancho_ocupacion_max_m": "Ancho máx. de ocupación (m)", "superficie_fondo_m2": "Superficie de fondo (m²)",
            "superficie_restauracion_m2": "Superficie a restaurar (m²)"}
        datos = [["Concepto", "Valor"]]
        for k, v in vol.items():
            if k in etiquetas and v is not None:
                datos.append([etiquetas[k], "Sí" if v is True else ("No" if v is False else _n(v, 2 if "cota" in k or "_m" == k[-2:] else 0))])
        s.append(_tabla(datos, [90 * mm, 84 * mm]))
        par = ", ".join(f"{k.replace('_', ' ')}: {('auto' if v is None else _n(v, 2))}" for k, v in r["parametros"].items())
        s.append(Paragraph(f"Parámetros: {par}.", peq))
        s.append(Paragraph("Volúmenes en banco calculados sobre la malla del MDT. Taludes en proporción H:V. "
                           "Estimación de prefactibilidad sin replanteo topográfico.", peq))
        if "corte" in planos:
            s.append(KeepTogether([_img(planos["corte"], 170 * mm, 140 * mm)]))
        if "perfil" in planos:
            s.append(_img(planos["perfil"], 170 * mm, 100 * mm))
            tr = r["vial"]["tramos"]
            datos = [["PK", "Desmonte (m³)", "Terraplén (m³)", "Pend. media (%)", "Altura máx. (m)"]]
            datos += [[f"{_n(x['pk_ini'])}–{_n(x['pk_fin'])}", _n(x["desmonte_m3"]), _n(x["terraplen_m3"]),
                       _n(x["pendiente_media_pct"], 1), _n(x["altura_max_m"], 1)] for x in tr[:60]]
            s.append(_tabla(datos, [34 * mm, 35 * mm, 35 * mm, 35 * mm, 35 * mm]))

    # ---------- 5. afecciones
    s.append(PageBreak())
    s.append(Paragraph("5. Condicionantes y afecciones", h2))
    datos = [["Capa", "Resultado", "Detalle"]]
    extra = []
    for i, af in enumerate(r["afecciones"], start=1):
        det = []
        for z in af.get("zonas", []):
            if z["afecta"]:
                el = f" ({', '.join(z.get('elementos', [])[:4])})" if z.get("elementos") else ""
                det.append(f"{z['etiqueta']}: {_n(z['area_m2'])} m² ({_n(z['pct'], 1)} %){el}")
        if not det:
            if af["estado"] == "sin_afeccion" and af.get("distancia_min_m") is not None:
                det.append(f"Elemento más cercano a {_n(af['distancia_min_m'])} m")
            elif af["estado"] in ("no_instalada", "sin_cobertura", "error"):
                det.append(ESTADO_TEXTO[af["estado"]])
        if af.get("cobertura_parcial") and af["estado"] != "afecta":
            det.append("<i>Cobertura parcial: resultado no concluyente</i>")
        datos.append([Paragraph(f"<b>{af['nombre']}</b><br/><font size=6.5 color='#666'>{af['tema']}</font>", celda),
                      Paragraph(f"<font color='white'><b>{SEM_TEXTO.get(af['semaforo'])}</b></font>", celda),
                      Paragraph("<br/>".join(det) or "—", celda)])
        extra.append(("BACKGROUND", (1, i), (1, i), SEM_COLOR.get(af["semaforo"], colors.grey)))
    s.append(_tabla(datos, [55 * mm, 25 * mm, 94 * mm], estilo_extra=extra))

    # ---------- 6. tramitación
    s.append(CondPageBreak(50 * mm))
    s.append(Paragraph("6. Orientación de tramitación", h2))
    datos = [["Prioridad", "Trámite", "Motivo"]]
    datos += [[t["nivel"].capitalize(), Paragraph(t["tramite"], celda), Paragraph(t["motivo"], celda)]
              for t in r["tramitacion"]]
    s.append(_tabla(datos, [20 * mm, 55 * mm, 99 * mm]))
    s.append(Paragraph("Orientación no vinculante. Debe confirmarse con los órganos competentes y la normativa "
                       "autonómica y municipal aplicable.", peq))

    # ---------- 7. fuentes y limitaciones
    s.append(CondPageBreak(60 * mm))
    s.append(Paragraph("7. Fuentes y limitaciones", h2))
    datos = [["Capa", "Fuente", "Versión", "Notas"]]
    for af in r["afecciones"]:
        datos.append([Paragraph(af["nombre"], celda), Paragraph(af["fuente"], celda), Paragraph(af["fecha"], celda),
                      Paragraph(af.get("notas", ""), celda)])
    if f:
        datos.append([Paragraph("Modelo del terreno", celda), Paragraph(f["fuente"], celda), "",
                      Paragraph(", ".join(f["ficheros"][:6]), celda)])
    s.append(_tabla(datos, [38 * mm, 45 * mm, 22 * mm, 69 * mm]))
    s.append(Spacer(1, 6))
    s.append(Paragraph(
        "Este documento es un estudio de prefactibilidad elaborado automáticamente a partir de cartografía oficial "
        "pública. No sustituye al proyecto técnico, al levantamiento topográfico, al estudio geotécnico ni a las "
        "consultas y autorizaciones de los organismos competentes. Las distancias de protección son aproximadas y "
        "las capas pueden estar incompletas o desactualizadas.", peq))

    doc.build(s, canvasmaker=lambda *a, **k: _Lienzo(*a, titulo=r["nombre"][:60], **k))


# ================================================================== API
def generar(id_: str) -> dict[str, Path]:
    """Genera (si no existen para la versión actual) PDF, Excel y ZIP."""
    d = carpeta(id_)
    r = resultado(id_)
    if r is None:
        raise FileNotFoundError("El estudio no tiene resultados.")
    dest = d / "informes"
    marca = dest / f"v{r['version']}.ok"
    base = f"Prefactibilidad_{_slug(r['nombre'])}_{r['id'][:8]}"
    rutas = {"pdf": dest / f"{base}.pdf", "xlsx": dest / f"{base}.xlsx", "zip": dest / f"{base}.zip"}
    if marca.exists() and all(p.exists() for p in rutas.values()):
        return rutas
    dest.mkdir(parents=True, exist_ok=True)
    planos = _planos(r, d, dest)
    excel(r, rutas["xlsx"])
    pdf(r, rutas["pdf"], planos)
    with zipfile.ZipFile(rutas["zip"], "w", zipfile.ZIP_DEFLATED) as z:
        z.write(rutas["pdf"], rutas["pdf"].name)
        z.write(rutas["xlsx"], rutas["xlsx"].name)
        for nombre, p in planos.items():
            z.write(p, f"planos/{p.name}")
        z.write(d / "ambito.geojson", "geojson/ambito.geojson")
        if (d / "visor/ocupacion.geojson").exists():
            z.write(d / "visor/ocupacion.geojson", "geojson/ocupacion.geojson")
        for f in sorted((d / "capas").glob("*.geojson")):
            z.write(f, f"geojson/afecciones/{f.name}")
        z.writestr("resultado.json", json.dumps(r, ensure_ascii=False, indent=1))
    marca.write_text("ok")
    return rutas


def _slug(t: str) -> str:
    import re
    import unicodedata
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_")[:40] or "estudio"
